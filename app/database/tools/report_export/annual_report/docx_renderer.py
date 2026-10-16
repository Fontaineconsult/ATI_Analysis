"""Render a report model to a Word document with python-docx.

Everything visual comes from theme.py. This module knows the document's structure:

    cover
    contents (a TOC field Word fills on open)
    1  Executive summary          counts, stat strip, status distribution by working group
    2..n  one section per working group
          one subsection per goal
          one page per indicator: header strip (status, review, priority, bar coverage),
          communities of practice, people, evidence summary with the administrative
          review beneath it, evidence table, plans, accomplishments, recommendations,
          concerns, notes, measures, systems
    Supplement  indicator bar coverage, one table per indicator that has a bar,
                linked from and back to the indicator
    Appendix A  status level definitions
    Appendix B  method

The renderer takes the model dict and never touches the graph.
"""
from __future__ import annotations

import re
from itertools import count
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

from . import theme as T

_bookmark_ids = count(1)
# The plain SFBRN wordmark, rendered once from app/static/img/sfbrn-logo.svg with its
# fills inlined (Word cannot take the SVG from python-docx). Regenerate the PNG if the
# SVG changes.
LOGO = Path(__file__).with_name("assets") / "sfbrn-logo.png"


# --------------------------------------------------------------------------- #
# Low-level helpers (python-docx has no API for these; they are plain OOXML)   #
# --------------------------------------------------------------------------- #
def _set_font(run, name=T.BODY_FONT, size=None, bold=None, italic=None, color=None, underline=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    if underline is not None:
        run.font.underline = underline


def _style_font(style, name, size, color=None, bold=None, italic=None):
    style.font.name = name
    style.font.size = size
    if color is not None:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    # The built-in heading and title styles point at the document theme's fonts,
    # and a theme attribute wins over an explicit face. Drop them or Georgia never
    # appears.
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def _type_label(type_name: str) -> str:
    """'InternalPolicy' -> 'Internal Policy'. Class names are not report vocabulary."""
    out = []
    for i, ch in enumerate(type_name or ""):
        if ch.isupper() and i and not type_name[i - 1].isupper():
            out.append(" ")
        out.append(ch)
    return "".join(out)


def _shade(cell, fill_hex: str):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcpr.append(shd)


def _cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblpr = table._element.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


def _table_borders(table, *, outer_hex=T.HAIRLINE, inner_hex=T.HAIRLINE, vertical=False,
                   top_size=8, bottom_size=8, inner_size=4):
    """Horizontal rules only: a top and bottom rule, hairlines between rows, no
    vertical lines. That is what makes a table read as a report table rather than
    a spreadsheet."""
    tblpr = table._element.tblPr
    borders = OxmlElement("w:tblBorders")
    spec = {
        "top": ("single", top_size, outer_hex),
        "bottom": ("single", bottom_size, outer_hex),
        "insideH": ("single", inner_size, inner_hex),
        "left": ("nil", 0, "auto"),
        "right": ("nil", 0, "auto"),
        "insideV": ("single" if vertical else "nil", inner_size if vertical else 0, inner_hex),
    }
    for edge, (val, size, color) in spec.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        if val != "nil":
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        borders.append(el)
    tblpr.append(borders)


def _paragraph_border_bottom(paragraph, color_hex: str, size=12, space=4):
    ppr = paragraph._element.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color_hex)
    pbdr.append(bottom)
    ppr.append(pbdr)


def _field(paragraph, instr: str, placeholder: str = ""):
    """A Word field (PAGE, NUMPAGES, TOC). Word computes it on open or update."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    run._element.append(begin)
    run = paragraph.add_run()
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    run._element.append(instr_el)
    run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._element.append(sep)
    run = paragraph.add_run(placeholder)
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._element.append(end)


def bookmark_name(prefix: str, key: str) -> str:
    """Word bookmark names allow letters, digits and underscores only."""
    return f"{prefix}_{re.sub(r'[^A-Za-z0-9]+', '_', key)}"


def _bookmark(paragraph, name: str):
    """Wrap the paragraph's content in a named bookmark so a hyperlink can target it."""
    bid = str(next(_bookmark_ids))
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bid)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bid)
    p = paragraph._element
    ppr = p.find(qn("w:pPr"))
    if ppr is not None:
        ppr.addnext(start)
    else:
        p.insert(0, start)
    p.append(end)


def _internal_link(paragraph, text: str, anchor: str, size=T.SIZE_BODY, bold=False):
    """A hyperlink to a bookmark in the same document, styled as a report link."""
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    link.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), T.BODY_FONT)
    rpr.append(rfonts)
    if bold:
        rpr.append(OxmlElement("w:b"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), T.STEEL_HEX)
    rpr.append(color)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size.pt * 2)))
    rpr.append(sz)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._element.append(link)


def _keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def _set_col_widths(table, widths: Iterable):
    widths = list(widths)
    for row in table.rows:
        for idx, w in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = w


def _mark_updatable_fields(document):
    """Ask Word to refresh fields (the TOC) when the file opens."""
    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


# --------------------------------------------------------------------------- #
# Styles                                                                       #
# --------------------------------------------------------------------------- #
def _install_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    _style_font(normal, T.BODY_FONT, T.SIZE_BODY, color=T.INK)
    normal.paragraph_format.space_after = T.SPACE_AFTER_BODY
    normal.paragraph_format.line_spacing = T.LINE_SPACING

    for name, size, color, before in (
        ("Heading 1", T.SIZE_H1, T.NAVY, T.SPACE_BEFORE_H1),
        ("Heading 2", T.SIZE_H2, T.NAVY, T.SPACE_BEFORE_H2),
        ("Heading 3", T.SIZE_H3, T.STEEL, T.SPACE_BEFORE_H3),
    ):
        st = styles[name]
        _style_font(st, T.HEADING_FONT, size, color=color, bold=False)
        st.paragraph_format.space_before = before
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.0

    h4 = styles["Heading 4"]
    _style_font(h4, T.BODY_FONT, T.SIZE_H4, color=T.NAVY, bold=True, italic=False)
    h4.paragraph_format.space_before = T.SPACE_BEFORE_H4
    h4.paragraph_format.space_after = Pt(3)
    h4.paragraph_format.keep_with_next = True

    title = styles["Title"]
    _style_font(title, T.HEADING_FONT, T.SIZE_TITLE, color=T.NAVY, bold=False)
    title.paragraph_format.space_after = Pt(4)

    for name, size, color, italic in (
        ("Report Kicker", T.SIZE_KICKER, T.STEEL, False),
        ("Report Meta", T.SIZE_SMALL, T.SLATE, False),
        ("Report Caption", T.SIZE_SMALL, T.SLATE, True),
        ("Report Table", T.SIZE_TABLE, T.INK, False),
        ("Report Quiet", T.SIZE_BODY, T.SLATE, True),
    ):
        st = styles.add_style(name, 1)  # WD_STYLE_TYPE.PARAGRAPH
        st.base_style = normal
        _style_font(st, T.BODY_FONT, size, color=color, italic=italic)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = 1.05

    styles["Report Kicker"].font.all_caps = True
    styles["Report Kicker"].font.bold = True
    styles["Report Kicker"].paragraph_format.space_before = Pt(8)
    styles["Report Table"].paragraph_format.space_after = Pt(0)
    styles["Report Table"].paragraph_format.line_spacing = 1.0

    stat = styles.add_style("Report Stat", 1)
    stat.base_style = normal
    _style_font(stat, T.HEADING_FONT, T.SIZE_STAT, color=T.NAVY)
    stat.paragraph_format.space_after = Pt(0)
    stat.paragraph_format.line_spacing = 1.0

    bullet = styles["List Bullet"]
    _style_font(bullet, T.BODY_FONT, T.SIZE_BODY, color=T.INK)
    bullet.paragraph_format.space_after = Pt(3)


def _page_setup(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = section.right_margin = T.PAGE_MARGIN
    section.top_margin = T.PAGE_MARGIN_TOP
    section.bottom_margin = T.PAGE_MARGIN
    section.header_distance = T.HEADER_DISTANCE


def _header_footer(section, model):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.text = ""
    run = p.add_run(f"{T.PROGRAM_NAME}  ·  {model['campus']['name']}  ·  {model['academic_year']}")
    _set_font(run, T.BODY_FONT, T.SIZE_SMALL, color=T.SLATE)
    _paragraph_border_bottom(p, T.HAIRLINE, size=4, space=2)

    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(f"{T.REPORT_TITLE}  ·  Page ")
    _set_font(run, T.BODY_FONT, T.SIZE_SMALL, color=T.SLATE)
    _field(p, "PAGE", "1")
    for r in p.runs[-4:]:
        _set_font(r, T.BODY_FONT, T.SIZE_SMALL, color=T.SLATE)


# --------------------------------------------------------------------------- #
# Building blocks                                                              #
# --------------------------------------------------------------------------- #
def _kicker(doc, text):
    p = doc.add_paragraph(text, style="Report Kicker")
    _keep_with_next(p)
    return p


def _body(doc, text, style=None):
    return doc.add_paragraph(text, style=style) if style else doc.add_paragraph(text)


def _quiet(doc, text):
    return doc.add_paragraph(text, style="Report Quiet")


def _meta(doc, text):
    return doc.add_paragraph(text, style="Report Meta")


def _label_value(paragraph, label, value, size=T.SIZE_BODY):
    r = paragraph.add_run(f"{label}  ")
    _set_font(r, T.BODY_FONT, size, bold=True, color=T.NAVY)
    r = paragraph.add_run(value)
    _set_font(r, T.BODY_FONT, size, color=T.INK)


def _table(doc, headers, rows, widths=None, *, header_fill=T.TINT, zebra=False,
           first_col_bold=False, font_size=T.SIZE_TABLE, vertical=False):
    """A report table: shaded header, horizontal rules, optional zebra."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = widths is None
    _table_borders(table, vertical=vertical)
    _cell_margins(table)

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.style = doc.styles["Report Table"]
        r = p.add_run(str(h))
        _set_font(r, T.BODY_FONT, font_size, bold=True, color=T.NAVY)
        _shade(cell, header_fill)

    for n, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.text = ""
            if zebra and n % 2 == 1:
                _shade(cell, T.TINT)
            _fill_cell(doc, cell, value, font_size, bold=(first_col_bold and i == 0))

    if widths:
        _set_col_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def _fill_cell(doc, cell, value, font_size, bold=False):
    """A cell value is a string, or a list of (text, kind) lines where kind is
    'strong' | 'normal' | 'quiet' | 'status:<level>' | 'link:<bookmark>'."""
    first = True
    lines = value if isinstance(value, list) else [(str(value), "strong" if bold else "normal")]
    for text, kind in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.style = doc.styles["Report Table"]
        if kind.startswith("status:"):
            level = kind.split(":", 1)[1] or "Unset"
            fill, color = T.STATUS_FILL.get(level, T.STATUS_FILL["Unset"])
            _shade(cell, fill)
            r = p.add_run(text or "Not set")
            _set_font(r, T.BODY_FONT, font_size, bold=True, color=color)
            continue
        if kind.startswith("link:"):
            _internal_link(p, text, kind.split(":", 1)[1], size=font_size)
            continue
        r = p.add_run(text)
        if kind == "strong":
            _set_font(r, T.BODY_FONT, font_size, bold=True, color=T.INK)
        elif kind == "quiet":
            _set_font(r, T.BODY_FONT, Pt(font_size.pt - 0.5), color=T.SLATE)
        else:
            _set_font(r, T.BODY_FONT, font_size, color=T.INK)


def _stat_strip(doc, stats: list[tuple[str, str]]):
    """Four big numbers in a row, the way a briefing opens."""
    table = doc.add_table(rows=2, cols=len(stats))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    _table_borders(table, outer_hex=T.NAVY_HEX, inner_hex=T.HAIRLINE, top_size=12, bottom_size=4, inner_size=0)
    _cell_margins(table, top=120, bottom=120, left=120, right=120)
    for i, (number, label) in enumerate(stats):
        top = table.rows[0].cells[i]
        top.text = ""
        p = top.paragraphs[0]
        p.style = doc.styles["Report Stat"]
        p.add_run(str(number))
        _shade(top, T.TINT_STRONG)
        bottom = table.rows[1].cells[i]
        bottom.text = ""
        p = bottom.paragraphs[0]
        p.style = doc.styles["Report Meta"]
        r = p.add_run(label)
        _set_font(r, T.BODY_FONT, T.SIZE_SMALL, color=T.SLATE)
        _shade(bottom, T.TINT_STRONG)
    doc.add_paragraph()


def _bullets(doc, items: Iterable[str]):
    for text in items:
        doc.add_paragraph(text, style="List Bullet")


def _dated_entries(doc, entries, *, name_key="name", date_key="date", body_key="content", author_key=None):
    for e in entries:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.2)
        # The date leads every entry. An entry the record never dated says so,
        # rather than looking like the date was dropped.
        head = " · ".join(x for x in (e.get(date_key) or "Undated", e.get(name_key)) if x)
        if author_key and e.get(author_key):
            head = f"{head} · {e[author_key]}"
        r = p.add_run(head)
        _set_font(r, T.BODY_FONT, T.SIZE_SMALL, bold=True, color=T.STEEL)
        p.add_run("\n")
        r = p.add_run(e.get(body_key, ""))
        _set_font(r, T.BODY_FONT, T.SIZE_SMALL, color=T.INK)


# --------------------------------------------------------------------------- #
# Sections                                                                     #
# --------------------------------------------------------------------------- #
def _cover(doc, model):
    if LOGO.is_file():
        doc.add_picture(str(LOGO), width=T.LOGO_WIDTH)
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Report Kicker")
    r = p.add_run(f"{T.SYSTEM_NAME}  ·  {T.PROGRAM_NAME}")
    _set_font(r, T.BODY_FONT, T.SIZE_KICKER, bold=True, color=T.STEEL)

    p = doc.add_paragraph(style="Title")
    p.add_run(T.REPORT_TITLE)
    _paragraph_border_bottom(p, T.NAVY_HEX, size=24, space=8)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run(model["campus"]["name"])
    _set_font(r, T.HEADING_FONT, T.SIZE_SUBTITLE, color=T.NAVY)
    p = doc.add_paragraph()
    r = p.add_run(f"Academic year {model['academic_year']}")
    _set_font(r, T.HEADING_FONT, T.SIZE_SUBTITLE, color=T.SLATE)

    for _ in range(14):
        doc.add_paragraph()
    s = model["summary"]
    _meta(doc, f"Prepared {model['generated']} from the ATI evidence record.")
    _meta(doc, f"{s['indicators']} success indicators across {s['working_groups']} priority areas.")
    _meta(doc, "Status levels follow the CSU ATI maturity model (Appendix A).")
    doc.add_page_break()


def _contents(doc):
    doc.add_paragraph("Contents", style="Heading 1")
    p = doc.add_paragraph()
    _field(p, 'TOC \\o "1-2" \\h \\z \\u', "Right-click and choose Update Field to build the table of contents.")
    doc.add_page_break()


def _executive_summary(doc, model):
    s = model["summary"]
    doc.add_paragraph("1  Executive summary", style="Heading 1")
    campus = model["campus"]["name"]
    year = model["academic_year"]
    reviewed_pct = round(100 * s["reviewed"] / s["indicators"]) if s["indicators"] else 0
    _body(doc, (
        f"This report records the state of {campus}'s Accessible Technology Initiative for the "
        f"{year} academic year. It covers {s['indicators']} success indicators across "
        f"{s['working_groups']} priority areas and {s['goals']} goals. Each indicator carries a "
        f"maturity status on the CSU six-level scale, the evidence the campus holds for it, the "
        f"people responsible, and the plans and recommendations that move it. "
        f"{s['reviewed']} indicators ({reviewed_pct} percent) have completed administrative review "
        f"this year, and {s['with_summary']} carry a written evidence summary."
    ))

    _stat_strip(doc, [
        (s["indicators"], "Success indicators reported"),
        (s["established_or_above"], "At Established or above"),
        (s["active_evidence_items"], "Active evidence items"),
        (s["open_recommendations"], "Open recommendations"),
    ])

    _kicker(doc, "Maturity status by priority area")
    levels = [lvl for lvl in s["status_order"] if any(
        s["status_distribution"][wg].get(lvl) for wg in s["status_distribution"])]
    headers = ["Priority area"] + levels + ["Total"]
    rows = []
    for wg in model["working_groups"]:
        counts = s["status_distribution"][wg["name"]]
        total = sum(counts.values())
        rows.append([wg["name"]] + [str(counts.get(lvl, 0) or "") for lvl in levels] + [str(total)])
    totals = ["All areas"]
    for lvl in levels:
        totals.append(str(sum(s["status_distribution"][wg].get(lvl, 0) for wg in s["status_distribution"]) or ""))
    totals.append(str(s["indicators"]))
    rows.append(totals)
    _table(doc, headers, rows, first_col_bold=True)
    _meta(doc, "Unset means the campus has not yet assigned a status for the year.")

    _kicker(doc, "What the record holds")
    _table(doc, ["Element", "Count"], [
        ["Evidence items listed (including retired)", str(s["evidence_items"])],
        ["Plans", str(s["plans"])],
        ["Accomplishments", str(s["accomplishments"])],
        ["Open recommendations", str(s["open_recommendations"])],
        ["Open concerns", str(s["open_concerns"])],
    ], widths=[Inches(4.5), Inches(1.5)])
    doc.add_page_break()


def _coverage_cell(ind):
    """The header's bar-coverage cell: a count that links to the supplement, or a
    quiet note when the indicator has no authored bar."""
    cov = ind["coverage"]
    if not cov["total"]:
        return [("No bar authored", "quiet")]
    return [(f"{cov['satisfied']} of {cov['total']} met", f"link:{bookmark_name('cov', ind['key'])}")]


def _status_strip(doc, ind):
    st = ind["status"]
    rv = ind["review"]
    fl = ind["flags"]
    review = "Complete" if rv["complete"] else ("Ready for review" if rv["ready"] else "Not reviewed")
    if rv["complete"] and rv["completed_date"]:
        review += f" · {rv['completed_date']}"
    if rv["complete"] and rv["completed_by"]:
        review += f" · {rv['completed_by']}"
    next_year = fl["will_work_on_next_year"]
    plan_flag = "Yes" if next_year else ("No" if next_year is False else "")
    rows = [[
        [(st["current"] or "Not set", f"status:{st['current'] or 'Unset'}")],
        [(st["previous"] or "Not set", f"status:{st['previous'] or 'Unset'}")] if st["previous"] else [("Not set", "quiet")],
        [(review, "normal")],
        [(fl["priority_level"] or "", "normal")],
        [(plan_flag, "normal")],
        _coverage_cell(ind),
    ]]
    _table(doc, ["Current status", "Prior year", "Administrative review", "Priority", "Planned next year", "Bar coverage"],
           rows, widths=[Inches(1.1), Inches(1.1), Inches(2.0), Inches(0.8), Inches(0.9), Inches(1.1)])


def _evidence_table(doc, ind):
    rows = []
    for im in ind["implementations"]:
        title_lines = [(im["title"] or "(untitled)", "strong")]
        if im["retired"]:
            note = "Retired" + (f" {im['retired_date']}" if im["retired_date"] else "")
            if im["retired_note"]:
                note += f": {im['retired_note']}"
            title_lines.append((note, "quiet"))
        if im["description"]:
            title_lines.append((im["description"], "normal"))
        if im["rationale"]:
            title_lines.append((f"Why it answers this indicator: {im['rationale']}", "quiet"))

        owner_lines = [(im["owner"] or "", "normal")]
        if im["owner_title"]:
            owner_lines.append((im["owner_title"], "quiet"))
        if im["accountable_working_group"]:
            owner_lines.append((f"Accountable: {im['accountable_working_group']}", "quiet"))
        for c in im["accountable_communities"]:
            owner_lines.append((c, "quiet"))
        if im["participants"]:
            owner_lines.append(("Participants: " + ", ".join(
                f"{p['name']} ({p['role']})" if p["role"] else p["name"] for p in im["participants"]
            ), "quiet"))

        strength = "" if im["strength"] in (None, "") else str(im["strength"])
        if im["control"]:
            strength = f"{strength} · {im['control']}".strip(" ·")

        docs = []
        for d in im["documents"]:
            # A document's address is a file store or a SharePoint share, which
            # nobody types from paper. The name is the citation.
            docs.append((d["name"], "normal"))
        for w in im["webpages"]:
            docs.append((w["name"], "normal"))
            if w["url"] and w["url"] != w["name"] and w["url"].lower().startswith("http"):
                docs.append((w["url"], "quiet"))
        if not docs:
            docs = [("No active documentation", "quiet")]

        rows.append([title_lines, [(_type_label(im["type"]), "normal")], owner_lines,
                     [(strength, "normal")], docs])
    _table(doc, ["Evidence", "Type", "Responsible", "Strength", "Documentation"], rows,
           widths=[Inches(2.6), Inches(0.85), Inches(1.35), Inches(0.6), Inches(1.6)])


def _people_bullets(doc, ind):
    _kicker(doc, "Responsible")
    if not ind["implementers"]:
        _quiet(doc, "No one is listed as responsible.")
        return
    items = []
    for person in ind["implementers"]:
        bits = person["name"]
        if person["title"]:
            bits += f", {person['title']}"
        roles = [r["name"] + (" (in position description)" if r["in_position_description"] else "")
                 for r in person["roles"] if r["name"]]
        if roles:
            bits += " · " + "; ".join(roles)
        items.append(bits)
    _bullets(doc, items)


def _community_bullets(doc, ind):
    if not ind["stakeholders"]:
        return
    _kicker(doc, "Communities of practice")
    _bullets(doc, [
        f"{c['name']} ({c['note']})" if c["note"] else c["name"] for c in ind["stakeholders"]
    ])


def _indicator(doc, ind, *, new_page=True):
    """One indicator per page. The first indicator of a goal shares its page with
    the goal heading, which carries the break instead, so no heading is stranded."""
    number = ind["number"] or ind["key"]
    h = doc.add_paragraph(f"{number}  {ind['text']}", style="Heading 3")
    h.paragraph_format.page_break_before = new_page
    _bookmark(h, bookmark_name("ind", ind["key"]))

    # Header: the strip, then the two lists that say who is around this work.
    _status_strip(doc, ind)
    _community_bullets(doc, ind)
    _people_bullets(doc, ind)

    # The review: the campus's own summary, with the administrative review beneath it.
    _kicker(doc, "Evidence summary")
    if ind["review"]["summary"]:
        _body(doc, ind["review"]["summary"])
    else:
        _quiet(doc, "No evidence summary has been written for this indicator this year.")
    if ind["review"]["notes"]:
        _kicker(doc, "Administrative review")
        _dated_entries(doc, ind["review"]["notes"], name_key="_none", author_key="author")

    _kicker(doc, "Evidence")
    if ind["implementations"]:
        _evidence_table(doc, ind)
    else:
        _quiet(doc, "No evidence is listed for this indicator this year.")

    if ind["plans"]:
        _kicker(doc, "Plans")
        rows = []
        for pl in ind["plans"]:
            lines = [(pl["name"], "strong")]
            if pl["description"]:
                lines.append((pl["description"], "normal"))
            if pl["abandoned"] and pl["abandoned_notes"]:
                lines.append((f"Abandoned: {pl['abandoned_notes']}", "quiet"))
            tags = []
            if pl["is_key_plan"]:
                tags.append("Key plan")
            if pl["is_campus_plan"]:
                tags.append("Campus plan")
            status = pl["status"] or ("Abandoned" if pl["abandoned"] else "")
            rows.append([lines, [(status, "normal")] + [(t, "quiet") for t in tags]])
        _table(doc, ["Plan", "Status"], rows, widths=[Inches(5.6), Inches(1.4)])

    if ind["accomplishments"]:
        _kicker(doc, "Accomplishments")
        _bullets(doc, [
            f"{a['name']}: {a['description']}" if a["description"] else a["name"]
            for a in ind["accomplishments"]
        ])

    if ind["recommendations"]:
        _kicker(doc, "Recommendations")
        rows = []
        for r in ind["recommendations"]:
            lines = [(r["text"], "strong")]
            if r["detail"]:
                lines.append((r["detail"], "normal"))
            if r["resolution"]:
                lines.append((f"Resolution: {r['resolution']}", "quiet"))
            meta = " · ".join(x for x in (r["date"], r["author"]) if x)
            rows.append([lines, [(r["status"] or "Open", "normal")] + ([(meta, "quiet")] if meta else [])])
        _table(doc, ["Recommendation", "Status"], rows, widths=[Inches(5.4), Inches(1.6)])

    if ind["concerns"]:
        _kicker(doc, "Concerns")
        rows = []
        for c in ind["concerns"]:
            lines = [(c["text"], "strong")]
            if c["detail"]:
                lines.append((c["detail"], "normal"))
            if c["became"]:
                lines.append((f"Became {c['became']}", "quiet"))
            meta = " · ".join(x for x in (c["date"], c["author"]) if x)
            rows.append([lines, [(c["status"] or "Open", "normal")] + ([(meta, "quiet")] if meta else [])])
        _table(doc, ["Concern", "Status"], rows, widths=[Inches(5.4), Inches(1.6)])

    if ind["notes"]:
        _kicker(doc, "Notes")
        _dated_entries(doc, ind["notes"])   # the model orders them newest first

    if ind["metrics"]:
        _kicker(doc, "Measures")
        rows = []
        for m in ind["metrics"]:
            lines = [(m["name"], "strong")]
            if m["description"]:
                lines.append((m["description"], "normal"))
            if m["comment"]:
                lines.append((m["comment"], "quiet"))
            rows.append([lines, [(m["value"], "normal")], [(m["type"], "normal")]])
        _table(doc, ["Measure", "Value", "Type"], rows, widths=[Inches(4.6), Inches(1.2), Inches(1.2)])

    systems = []
    if ind["assets"]:
        systems.append(("Systems", ", ".join(
            f"{a['title']} ({a['scope']})" if a["scope"] else a["title"] for a in ind["assets"])))
    if ind["interfaces"]:
        systems.append(("Interfaces", ", ".join(
            f"{i['title']} ({i['function']})" if i["function"] else i["title"] for i in ind["interfaces"])))
    if ind["tools"]:
        systems.append(("Tools", ", ".join(ind["tools"])))
    if ind["vendors"]:
        systems.append(("Vendors", ", ".join(ind["vendors"])))
    if ind["taaps"]:
        systems.append(("Timelines and action plans", ", ".join(t["title"] for t in ind["taaps"])))
    if systems:
        _kicker(doc, "Systems and tools in scope")
        for label, value in systems:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            _label_value(p, label, value, size=T.SIZE_SMALL)


def _working_group(doc, wg, index, model):
    h = doc.add_paragraph(f"{index}  {wg['name']}", style="Heading 1")
    h.paragraph_format.page_break_before = True
    s = model["summary"]["status_distribution"][wg["name"]]
    n = sum(s.values())
    top = [lvl for lvl in model["summary"]["status_order"] if s.get(lvl)]
    _body(doc, (
        f"{wg['name']} carries {n} success indicators across {len(wg['goals'])} goals this year. "
        + ("Status by level: " + ", ".join(f"{s[lvl]} {lvl}" for lvl in top) + "." if top else "")
    ))
    _kicker(doc, "Goals")
    _table(doc, ["Goal", "Indicators"], [
        [f"Goal {g['number']}: {g['name']}", str(len(g["indicators"]))] for g in wg["goals"]
    ], widths=[Inches(6.0), Inches(1.0)], first_col_bold=False)

    for g in wg["goals"]:
        gh = doc.add_paragraph(f"Goal {g['number']}  {g['name']}", style="Heading 2")
        gh.paragraph_format.page_break_before = True
        for n, ind in enumerate(g["indicators"]):
            _indicator(doc, ind, new_page=(n > 0))


def _supplement(doc, model, index):
    """Bar coverage, indicator by indicator, out of the body so each indicator's page
    stays readable. Every entry is bookmarked for the header's link and links back."""
    h = doc.add_paragraph(f"{index}  Supplement: indicator bar coverage", style="Heading 1")
    h.paragraph_format.page_break_before = True
    _body(doc, (
        "Some indicators carry a companion bar: the specific things the indicator asks for, "
        "one requirement at a time, each marked with the evidence that claims to answer it. "
        "A claim is an assertion by the campus, not a finding. The count in each indicator's "
        "header links to its table here, and each table links back."
    ))
    without = []
    for wg in model["working_groups"]:
        for g in wg["goals"]:
            for ind in g["indicators"]:
                cov = ind["coverage"]
                if not cov["total"]:
                    without.append(ind["number"] or ind["key"])
                    continue
                number = ind["number"] or ind["key"]
                head = doc.add_paragraph(f"{number}  {ind['text']}", style="Heading 3")
                _bookmark(head, bookmark_name("cov", ind["key"]))
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                r = p.add_run(f"{wg['name']} · Goal {g['number']} · {cov['satisfied']} of {cov['total']} requirements met · ")
                _set_font(r, T.BODY_FONT, T.SIZE_SMALL, color=T.SLATE)
                _internal_link(p, "Return to the indicator", bookmark_name("ind", ind["key"]), size=T.SIZE_SMALL)
                rows = []
                for req in cov["requirements"]:
                    if req["satisfied"]:
                        state = "Met"
                    elif not req["implementation_evidenced"]:
                        state = "Not evidenced by an implementation"
                    else:
                        state = "Not met"
                    by = ", ".join(req["satisfied_by"])
                    rows.append([
                        [(req["requirement"], "normal")]
                        + ([(f"{req['level'].capitalize()} · {req['element']}", "quiet")] if req["level"] or req["element"] else []),
                        [(state, "strong" if req["satisfied"] else "normal")] + ([(by, "quiet")] if by else []),
                    ])
                _table(doc, ["Requirement", "State"], rows, widths=[Inches(4.8), Inches(2.2)])
    if without:
        _kicker(doc, "Indicators without an authored bar")
        _body(doc, ", ".join(without) + ".")


def _appendix_rubric(doc, model):
    h = doc.add_paragraph("Appendix A  Status level definitions", style="Heading 1")
    h.paragraph_format.page_break_before = True
    _body(doc, "The CSU Accessible Technology Initiative grades each success indicator on a "
               "six-level maturity scale. A level is earned only when procedures, resources and "
               "documentation all meet its bar.")
    rows = []
    for r in model["rubric"]:
        rows.append([
            [(r["level"], f"status:{r['level']}")],
            [(r["procedures"], "normal")],
            [(r["resources"], "normal")],
            [(r["documentation"], "normal")],
        ])
    _table(doc, ["Level", "Procedures", "Resources", "Documentation"], rows,
           widths=[Inches(1.0), Inches(2.0), Inches(2.0), Inches(2.0)], vertical=True)


def _appendix_method(doc, model):
    h = doc.add_paragraph("Appendix B  Method", style="Heading 1")
    h.paragraph_format.page_break_before = True
    _bullets(doc, [
        f"Generated {model['generated']} from the ATI evidence record for {model['campus']['name']}, "
        f"academic year {model['academic_year']}. The content matches the dashboard's report view for "
        "each indicator on that date.",
        "Every note, measure, document, webpage and plan is included unless it was marked not for "
        "the report. A documentation link curated into or out of this year overrides that mark. "
        "Correspondence is not reproduced.",
        "Retired evidence stays listed and is marked retired, because the evidence chain is historical.",
        "Documents marked out of date and web pages that no longer exist are omitted.",
        "Indicators removed from the CSU indicator set are omitted.",
        "Strength is the curator's rating, from 0 to 3, of how directly an evidence item answers the "
        "indicator. The control mark says whether the campus operates the practice itself (internal) "
        "or relies on another unit (external).",
        "Bar coverage counts the requirements of an indicator's companion bar that listed evidence "
        "claims to meet. The tables are in the Supplement.",
        "Status levels are assigned by the campus through administrative review and are not changed "
        "by this report.",
    ])


# --------------------------------------------------------------------------- #
# Entry                                                                        #
# --------------------------------------------------------------------------- #
def render(model: dict, path: str) -> str:
    """Write the document for ``model`` to ``path`` and return the path."""
    global _bookmark_ids
    _bookmark_ids = count(1)   # ids restart per document, so equal models give equal files
    doc = Document()
    _install_styles(doc)
    _page_setup(doc)
    _cover(doc, model)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.left_margin = section.right_margin = T.PAGE_MARGIN
    section.top_margin = T.PAGE_MARGIN_TOP
    section.bottom_margin = T.PAGE_MARGIN
    _header_footer(section, model)
    # Restart page numbering after the cover.
    sectpr = section._sectPr
    pgnum = OxmlElement("w:pgNumType")
    pgnum.set(qn("w:start"), "1")
    sectpr.append(pgnum)

    _contents(doc)
    _executive_summary(doc, model)
    index = 2
    for wg in model["working_groups"]:
        _working_group(doc, wg, index, model)
        index += 1
    _supplement(doc, model, index)
    _appendix_rubric(doc, model)
    _appendix_method(doc, model)

    _mark_updatable_fields(doc)
    doc.save(path)
    return path
