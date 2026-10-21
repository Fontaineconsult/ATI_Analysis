"""Build SFBRN-Accessible-Procurement-Rubric.docx.

The rubric prose is the Definitions section of review-matrix-tester.html, rendered from
review-matrix.js, so the document carries the same numbers as the model. Prose follows app/database/ontology/writing-style.md.

The rules, principles and calibration choices behind each section are Word endnotes, marked on the
section's first paragraph and collected under the closing Notes heading.

Document rule: no bold inside a paragraph. The source marks key terms with <strong> and
**bold** for the web page; this document renders them as plain text. Bold is kept only for
the label that opens a list entry or a paragraph, which the writing style's reference-document
rule asks for.

    .venv/Scripts/python app/database/ontology/impact-matrix/build-rubric-docx.py

Needs python-docx, pywin32 and Word (which fills the contents page), and the Playwright
install under e2e/node_modules.
"""
import json
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import date
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]
OUT = HERE / 'SFBRN-Accessible-Procurement-Rubric.docx'
PLAYWRIGHT = ROOT / 'e2e' / 'node_modules' / 'playwright'

# ---------------------------------------------------------------------------------------
# Grounding, by Definitions heading. Only rules that bind the CSU count as rules: federal
# and California law, federal standards California applies, and CSU policy. A principle is
# a commitment the framework holds (a Principle node in the graph, or one proposed in
# matrix-grounding-review.md). Calibration names the choices no rule sets.
# ---------------------------------------------------------------------------------------
GROUNDING = {
    'Accessible Procurement Process Overview': {
        'rules': [
            ('CSU ATI Procurement Process, Step 3', 'The ATI reviewer approves or denies the acquisition on the review findings. A denial is recorded with a detailed rationale, and each campus provides a way to request reconsideration.'),
            ('28 CFR 35.200', 'Web content and mobile apps that a public university provides, directly or through a vendor, must conform to WCAG 2.1 AA.'),
            ('Section 508 Standards, E202.7 (36 CFR 1194), through Gov Code 7405', 'When no conforming product meets the need, the CSU buys the product that best meets the standard, documents why, and provides the information and functions by an alternative means.'),
            ('CSU TAAP Authoring Template, instructions', 'A TAAP is needed when barriers are "severe enough to prevent users from performing essential functions". Failures "that do not prohibit core functionality do not require this process."'),
        ],
        'principles': [
            ('Time-bound alternative access when full conformance is not yet achieved', 'Where content cannot yet meet the standard, the institution provides a documented, time-limited, equally effective alternative while it remediates, and does not treat the non-conformance as acceptable. A TAAP is that alternative, never compliance.'),
            ('Deference is earned through a documented, individualized process', 'Courts and OCR defer to an institution\'s judgment only when a record shows the responsible officials considered the individual and the alternatives. Every hold names its holder and every decision its reasons for that record.'),
        ],
        'calibration': 'The split into risk grade, holds and remedy step is the matrix\'s own design. No rule prescribes a hold state, a holder or a resume point.',
    },
    'Review standard': {
        'rules': [
            ('28 CFR 35.200', 'Web content and mobile apps that a public university provides, directly or through a vendor, must conform to WCAG 2.1 AA.'),
            ('Gov Code 7405 and 11135; Section 508 Standards (36 CFR 1194)', 'California applies the Section 508 Standards to the technology state entities buy, which sets the standard for software and hardware that is not web or mobile.'),
        ],
        'principles': [],
        'calibration': 'None. The kind of technology selects the standard the law names for it.',
    },
    'Essential function': {
        'rules': [
            ('WCAG 2.1, conformance requirement 3', 'Every page in a complete process must conform, so an essential function is judged across its whole process.'),
            ('CSU ATI Procurement Process, Step 1', 'The requester states the product\'s "essential functionality", which is the source of the essential function list.'),
        ],
        'principles': [],
        'calibration': 'None. The test is WCAG-EM 2.0\'s definition of essential functionality, and the source of the list is CSU Procurement Step 1.',
    },
    'Nonconformance and barrier': {
        'rules': [
            ('Section 508 Standards, Chapter 3 (36 CFR 1194)', 'The functional performance criteria name the nine disability groups a barrier is recorded against.'),
        ],
        'principles': [
            ('Severity is judged by the effect on an essential function', 'A barrier\'s severity comes from what it does to a named group of users on an essential function. Neither an ACR\'s conformance label nor a criterion\'s level says how severe a barrier is.'),
        ],
        'calibration': 'None. The split between a failure in the product and its effect on a person follows the law, which measures the effect.',
    },
    'Minimal impact': {
        'rules': [
            ('28 CFR 35.205', 'A nonconformance counts as compliant only when its impact is so minimal that people with disabilities still get the same information, interactions, transactions and participation, with substantially equivalent timeliness, privacy, independence and ease of use. The public entity must demonstrate it. The test measures the effect on a person, which is why severity attaches to the barrier and not to the failed criterion.'),
        ],
        'principles': [],
        'calibration': 'The four tests are the regulation\'s. Applying them to non-web products, where the Section 508 Standards have no minimal-impact clause, is the matrix\'s choice.',
    },
    'Finding': {
        'rules': [],
        'principles': [
            ('The audience determines which duty governs', 'Who uses a system decides which law governs its accessibility, so a finding records the role.'),
        ],
        'calibration': 'None. The four fields are the dimensions the law measures: who is affected, on which function, and how badly.',
    },
    'Exposure': {
        'rules': [
            ('ADA Title II, Section 504 and CSU policy', 'No binding rule contains a user-count, criticality or exposure threshold. Exposure therefore sets how deep the review goes, how strong the evidence must be and who signs, and never whether access is owed.'),
            ('CSU ATI Procurement Process', 'The depth of review depends on how the product will be used, which is why unknown answers hold the review.'),
            ('28 CFR 35.205 and Section 504 (34 CFR 104)', 'The transactions test and the duty across academic programs ground the floors for coursework and for the only route to a service.'),
            ('28 CFR 35.200', 'Content the university provides to the public falls under the web rule, which grounds the public floor.'),
            ('ADA Title I (29 CFR 1630) and FEHA', 'Employment and hiring carry their own duty, which grounds the job-duty floor.'),
            ('CSU Executive Order 1111', 'A known user with a disability or an open accommodation request makes the barrier concrete, which grounds that floor.'),
        ],
        'principles': [
            ('Exposure sets scrutiny, not entitlement (proposed)', 'How widely a product is used decides how closely it is reviewed and who signs. It does not decide whether people with disabilities are owed access.'),
            ('Prioritization by impact under finite resources', 'Campuses direct accessibility work at the barriers with the greatest impact rather than spreading it evenly.'),
            ('Same facts, same answer (proposed)', 'Two reviewers given the same facts reach the same result. Each question maps to one group so no fact is counted twice.'),
        ],
        'calibration': 'The user count cuts (25 and 250), the band cut points (0.25 and 0.55), the equal group weights and the assignment of each floor to medium or high are pilot calibration. No rule sets them.',
    },
    'Severity': {
        'rules': [
            ('28 CFR 35.205', 'The four tests define S1. A barrier that makes an essential function clearly harder to perform fails the minimal-impact test, so S2 does not conform. The university must be able to show minimal impact, so S1 needs demonstration.'),
            ('CSU TAAP Authoring Template, instructions', 'Barriers "severe enough to prevent users from performing essential functions" require a TAAP, which is the S3 line.'),
            ('Section 508 Standards, E202.7 (36 CFR 1194)', 'Section 508 has no minimal-impact exception, so for non-web technology an S1 or S2 product is acceptable only when no conforming product meets the need.'),
        ],
        'principles': [
            ('Severity is judged by the effect on an essential function','Severity derives from the impact on an essential function, independent of the vendor label and the criterion level.'),
            ('Severity calls need anchors', 'Trained reviewers disagree on severity often enough that one reviewer\'s call carries known error, so each level has a written anchor and a call made from documents alone stays open to testing.'),
            ('Remedy proportionate to barrier severity', 'Formal alternative-access planning applies to barriers that block essential functions, not to every conformance flaw.'),
        ],
        'calibration': 'The accumulation thresholds (3 S2 findings stacked on one function and group, and 6 pairs at high exposure) are the matrix\'s. The reasoning is that a user performs a function only by passing every barrier on it, so stacked S2 barriers block the function, and S2 barriers across most of the product exceed minimal impact in aggregate. The GSA labels minor, major and critical are used so vendor reports read across, and they are not a rule.',
    },
    'Evidence (Veracity)': {
        'rules': [
            ('CSU ATI Procurement Process, Step 2', 'The CSU asks for a current Accessibility Conformance Report for each interface before it buys, which is hold 1.'),
            ('CSU ATI success indicator 1.8-pro', 'Each campus has a process to validate vendors\' conformance claims, which is the loop under hold 3.'),
            ('CSU ATI Procurement Process, Step 3', 'The ATI reviewer decides on the review findings, so the findings are confirmed before the decision is made.'),
            ('28 CFR 35.205', 'The university must be able to show that a nonconformance has minimal impact, which is why S1 always needs a demonstration, an audit or ATI testing.'),
            ('Section 508 Standards, E202.7.2 (36 CFR 1194); CSU TAAP Authoring Template', 'Where conforming technology is not available, access is provided by an alternative means planned in advance. The fallback applies that remedy when conformance cannot be shown.'),
        ],
        'principles': [
            ('Vendor conformance claims require proportionate validation', 'A vendor\'s conformance report is an unverified claim. The institution checks it in proportion to the product\'s exposure, and a claim the product contradicts counts as no evidence.'),
            ('Report quality and product accessibility are separate measures', 'How far an ACR can be trusted and how accessible the product is are two measures, scored apart. A report with no reported failures is a warning sign.'),
            ('Program accessibility is a proactive, systemic duty', 'Access is planned before deployment, whether or not anyone has asked. Unproven conformance is treated as hiding a blocking barrier for that reason.'),
        ],
        'calibration': 'The floor values (any evidence at low exposure, E1 at medium, E2 at high), the two-year ACR age and the eight ACR quality checks are calibration, drawn from CSU practice rather than a rule. The reasoning that a Deny stands on weak evidence is the matrix\'s own. Vendor self-reports understate barriers and do not invent them, so a standard computed from the findings on file is a minimum.',
    },
    'Breadth': {
        'rules': [
            ('28 CFR 35.130(b)(3)', 'A university may not use methods of administration that have the effect of discriminating on the basis of disability. A product whose barriers reach most essential functions and groups has that effect more widely, which is what Breadth measures.'),
            ('CSU TAAP Authoring Template, Process Outcome and Institutional Risk', 'The template grades a plan by how many requirements it meets. Breadth raises the pressure, and pressure raises how many a plan must meet.'),
        ],
        'principles': [
            ('Remedy proportionate to barrier severity', 'The weight of the remedy matches the severity of the barrier, and Breadth is the part of severity that measures reach.'),
        ],
        'calibration': 'The blocked-pair steps (2 and 4), the S2 count (3) and the cap at 2 are the matrix\'s. The assumption of Breadth 2 under the fallback with no evidence at all mirrors how unknown exposure answers map to high.',
    },
    'TAAP': {
        'rules': [
            ('CSU TAAP Authoring Template, Requirements Checklist', 'To qualify as an equally effective alternative, the plan must meet all six requirements. The six listed are the template\'s own words.'),
            ('CSU TAAP Authoring Template, Product Specific Accessibility Statement, Administrative Approval and Annual Review', 'The TAAP includes a product-specific statement posted wherever the product is in use, is signed by the Department Chair or Manager, the Dean or Division Vice President and the ADA Compliance Officer, and "must be reviewed and updated annually or when renewing the ICT".'),
            ('CSU TAAP Authoring Template, If the Vendor Fails to Meet Compliance Requirements', 'A TAAP stays valid only while the vendor is actively remediating. Without measurable progress, keeping the product needs a formal risk memo signed by executive leadership.'),
            ('CSU ATI Procurement Process, Step 3; Section 508 Standards, E202.7.2 (36 CFR 1194)', 'The TAAP provides equally effective alternate access "in a timely manner, for anyone that later reports a problem accessing the product". The duty is the institution\'s and is planned before anyone asks.'),
            ('28 CFR 35.130(b)(1); 28 CFR 35.160(b)(2); 28 CFR 35.130(e); 34 CFR 104.4(b)', 'An alternative must be as effective as what others get, timely, private and usable independently, and a person may decline it.'),
            ('CSU Executive Order 1111', 'An accommodation comes from an interactive process with one person, in response to that person\'s request. TAAP support is given on the terms the plan sets, to anyone who reports a problem, and the review never relies on accommodation.'),
            ('28 CFR 35.164 and 35.204; Section 508 Standards, E202.6', 'Only the head of the university or a designee may decide, in writing and after considering all resources available to the program, that conforming is an undue burden or a fundamental alteration. The university must still give access to the maximum extent possible. Because the decision weighs all resources of the program, the cost of one purchase on its own does not establish the burden.'),
        ],
        'principles': [
            ('Equally effective access', 'Accessible technology must give a person with a disability the same result, benefit or opportunity as others, not merely some alternative.'),
            ('Undue burden and fundamental alteration bound the duty', 'The duty has limits, but the institution must still take other action that gives people with disabilities the program\'s benefits to the maximum extent possible. That is why the exception is claimed against the TAAP standard and never at intake.'),
            ('Notice is part of access, not administration around it', 'A remedy the affected person cannot find is not a remedy, which is why the TAAP statement is posted where the product is used.'),
            ('Conformance and accommodation are parallel duties', 'Approval under a TAAP never ends the individual accommodation duty.'),
        ],
        'calibration': 'None. Every term here is the template\'s or the law\'s.',
    },
    'Risk grade': {
        'rules': [
            ('CSU TAAP Authoring Template, instructions', 'A TAAP is needed only for barriers that prevent essential functions, so S0 and S1 are No alternate access required and S3 is Alternate access required at every exposure.'),
            ('28 CFR 35.205; CSU Executive Order 1111, section IV', 'At medium or high exposure, barriers that make essential functions clearly harder to perform reach too many people for the impact to be minimal, so S2 is Alternate access required there.'),
            ('Section 508 Standards, E202.7 (36 CFR 1194)', 'For non-web technology, a high-exposure S3 approval also carries a best-meets record with market research.'),
        ],
        'principles': [
            ('Remedy proportionate to barrier severity', 'The grade places the formal remedy on the barriers that block essential functions.'),
            ('Exposure sets scrutiny, not entitlement (proposed)', 'Exposure moves the S2 row only, and never makes an S3 product acceptable.'),
        ],
        'calibration': 'The No alternate access required grade for S2 at low exposure is the matrix\'s choice. The law does not scale with exposure, and the product still fails the standard, which the approval says in writing.',
    },
    'Pressure and the TAAP standard': {
        'rules': [
            ('CSU TAAP Authoring Template, Process Outcome and Institutional Risk', 'A plan meeting all six requirements is Equally Effective with low institutional risk, one to five is a Non-Equal Alternative with moderate risk, and none is a need for individualized accommodation with high risk.'),
            ('CSU TAAP Authoring Template, Administrative Approval', 'Each signer confirms they have "evaluated the associated institutional risks" and found the plan "an acceptable interim solution." A plan below the standard is a risk the signers accept.'),
            ('28 CFR 35.130(b)(1); 34 CFR 104.4(b)', 'An alternative with fewer functions is not as effective as what others get. The law does not scale with exposure, so a plan below six requirements is an accepted risk and never a compliant result.'),
        ],
        'principles': [
            ('Remedy proportionate to barrier severity', 'The number of requirements a plan must meet rises with the reach of the barriers and the exposure of the product.'),
        ],
        'calibration': 'The counts (3 of 6 at pressure 0 or 1, 4 at 2, all 6 at 3 or more) are the matrix\'s. The template sets the risk tiers, and no rule sets how many requirements a plan must meet at each exposure. A lower count at low pressure is deliberate, because TAAP work costs staff time, and the standard binds at every pressure.',
    },
    'Warning level': {
        'rules': [
            ('CSU TAAP Authoring Template, Administrative Approval', 'Each signer confirms they have "evaluated the associated institutional risks and accommodation requirements" and found the plan "an acceptable interim solution." The warning level tells the signers, by consequence, what they accept.'),
            ('CSU TAAP Authoring Template, Institutional Risk', 'The template grades a plan by the count of the six requirements it meets: all six low, one to five moderate, none high. That count is the gap term of the warning. The template\'s tiers grade the plan alone, so a full plan the template calls low risk can carry an elevated or high warning here from the audience it must serve.'),
            ('CSU ATI Procurement Process, Step 2', 'The CSU names the consequences that make a purchase high impact: denied access, a critical program, high accommodation cost, significant legal exposure and no known workaround. The warning sentence states residual risk in those terms.'),
        ],
        'principles': [
            ('Residual risk is stated by consequence and owned by a named officer (proposed)', 'The record states who cannot do what without staff support, what the department carries and which officer accepted that. Each level has a fixed meaning and action, the numbers print with the word, and the level maps to the TAAP signers.'),
            ('Time-bound alternative access when full conformance is not yet achieved', 'The acceptance holds only while the vendor shows progress, so the warning is reviewed at the first roadmap milestone.'),
        ],
        'calibration': 'The pressure term and the cut points (elevated at 2, high at 4) are the matrix\'s. No rule scales residual risk by audience. The nearest precedent is the Section508.gov risk model, which sums scope of use with the coverage of the alternate means plan and the remediation time. The level never lowers the TAAP standard.',
    },
    'Final decision': {
        'rules': [
            ('28 CFR 35.200; Gov Code 7405 and 11135', 'Web and mobile content must conform to WCAG 2.1 AA, and California applies the Section 508 Standards to the technology the CSU buys.'),
            ('CSU ATI Procurement Process', 'The CSU buys the most accessible product that meets the business need, and the ATI reviewer approves or denies on the findings, with a recorded rationale and a reconsideration path.'),
            ('CSU Accessible Technology Initiative Policy (2024)', '"The CSU seeks to improve product accessibility through partnerships and by leveraging the procurement process."'),
            ('28 CFR 35.130(b)(1) and (b)(3); 34 CFR 104.4(b)(1)', 'A university may not give people with disabilities a less effective service, or use methods of administration that have a discriminating effect. Requiring a product some users cannot use can have that effect.'),
            ('Section 508 Standards, E202.7 (36 CFR 1194)', 'When a more accessible product meets the need, the CSU chooses it.'),
            ('CSU TAAP Authoring Template, Process Outcome; CSU Executive Order 1111, section IV.C.4', 'A plan that meets none of the six requirements is "unable to provide alternative means of access". Access would then rest on individual accommodation, which the CSU\'s Universal Design commitment aims to reduce.'),
            ('CSU Executive Order 1111, sections V.B and VI.B; 28 CFR 35.130(b)(7); 29 CFR 1630.2(o)', 'An accommodation comes from the interactive process between one person and the campus. A purchase that carries it out is approved at intake, and the review of the product for a general audience does not apply.'),
        ],
        'principles': [
            ('Procurement as an accessibility lever', 'The institution advances accessibility at the point of purchase. A denial states its findings, the standard it applied and the route to a new review.'),
            ('Universal Design reduces accommodation', 'Designing for accessibility up front reduces the need for individual accommodations, so the review prefers the more accessible product.'),
        ],
        'calibration': 'None. Each deny condition rests on a rule above.',
    },
    'Holds': {
        'rules': [
            ('CSU ATI Procurement Process, Step 2', 'A current ACR for each interface before purchase (hold 1).'),
            ('CSU ATI Procurement Process', 'The depth of review depends on how the product will be used (hold 2).'),
            ('CSU ATI success indicator 1.8-pro; CSU ATI Procurement Process, Step 3', 'A process to validate vendor claims, and a decision made on confirmed findings (hold 3).'),
            ('CSU TAAP Authoring Template', 'The TAAP standard a plan is held to before the remedy step resumes (hold 4).'),
            ('28 CFR 35.164 and 35.204; Section 508 Standards, E202.6', 'The written determination belongs to the president or a designee, after considering all resources (hold 5).'),
        ],
        'principles': [
            ('Deference is earned through a documented, individualized process', 'Each hold names who holds it and what releases it, so the record shows the review waited for a fact rather than assumed one.'),
            ('Undue burden and fundamental alteration bound the duty', 'The claim leaves the reviewer and goes to the officer the law names.'),
        ],
        'calibration': 'The numbering, the holders and the resume points are the matrix\'s design.',
    },
}


# ---------------------------------------------------------------------------------------
# HTML to docx. The Definitions markup is small: h3, p, ul/li, table, strong, code.
# ---------------------------------------------------------------------------------------
class Defs(HTMLParser):
    """Collects the Definitions markup as a list of blocks: ('h3', runs), ('p', runs),
    ('li', runs), ('table', [[runs, ...], ...])."""

    def __init__(self):
        super().__init__()
        self.blocks = []
        self.runs = None          # current paragraph runs: [(text, bold)]
        self.bold = 0
        self.table = None         # rows of cells, each cell a run list
        self.cell = None
        self.kind = None

    def handle_starttag(self, tag, attrs):
        if tag in ('h3', 'p', 'li'):
            self.kind, self.runs = tag, []
        elif tag == 'table':
            self.table = []
        elif tag == 'tr':
            self.table.append([])
        elif tag in ('td', 'th'):
            self.cell = []
            self.runs = self.cell
        elif tag == 'strong':
            self.bold += 1

    def handle_endtag(self, tag):
        if tag in ('h3', 'p', 'li'):
            self.blocks.append((tag, self.runs))
            self.runs = None
        elif tag in ('td', 'th'):
            self.table[-1].append(self.cell)
            self.cell = self.runs = None
        elif tag == 'table':
            self.blocks.append(('table', self.table))
            self.table = None
        elif tag == 'strong':
            self.bold -= 1

    def handle_data(self, data):
        if self.runs is not None and data:
            self.runs.append((data, self.bold > 0))


def add_runs(par, runs):
    for text, bold in runs:
        r = par.add_run(text)
        r.bold = bool(bold)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def add_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
    t.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            par = c.paragraphs[0]
            par.paragraph_format.space_before = Pt(2)
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, [(txt, bold or i == 0) for txt, bold in cell])
            for r in par.runs:
                r.font.size = Pt(10)
            if i == 0:
                shade(c, HEADER_FILL)
    indent_table(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def indent_table(t):
    """Moves the table to the body column and fits it to the width that remains."""
    tbl_pr = t._tbl.tblPr
    ind = OxmlElement('w:tblInd')
    ind.set(qn('w:w'), str(int(BODY_INDENT.inches * 1440)))
    ind.set(qn('w:type'), 'dxa')
    tbl_pr.insert_element_before(ind, 'w:tblBorders', 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook', 'w:tblCaption', 'w:tblDescription')
    width = tbl_pr.find(qn('w:tblW'))
    if width is None:
        width = OxmlElement('w:tblW')
        tbl_pr.append(width)
    width.set(qn('w:w'), str(int((TEXT_WIDTH.inches - BODY_INDENT.inches) * 1440)))
    width.set(qn('w:type'), 'dxa')


def field(par, instr, placeholder=''):
    """A Word field (PAGE, TOC) in the paragraph. Word fills it on update."""
    run = par.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    text = OxmlElement('w:instrText'); text.set(qn('xml:space'), 'preserve'); text.text = instr
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    shown = OxmlElement('w:t'); shown.text = placeholder
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for el in (begin, text, sep, shown, end):
        run._r.append(el)


# Endnotes. python-docx cannot write Word endnotes, so the builder leaves a marker token at
# each citation point and Word replaces every token with a real endnote (Endnotes.Add) when
# it opens the file to fill the contents page. ENDNOTES holds (token, text) in document order.
ENDNOTES = []
KIND = {'rules': ('Binding rule', 'Binding rules'), 'principles': ('Principle', 'Principles'), 'calibration': ('Calibration', 'Calibration')}


def mark(par, heading):
    """Appends the section's endnote markers to the paragraph: one note per kind (binding rules,
    principles, calibration), each entry on its own line inside the note, with a superscript
    comma between markers so Word renders 1,2,3 rather than 123."""
    g = GROUNDING.get(heading)
    if not g:
        return
    groups = [(key, g[key]) for key in ('rules', 'principles') if g[key]]
    if not g['calibration'].startswith('None.'):
        groups.append(('calibration', [(None, g['calibration'])]))
    for i, (key, entries) in enumerate(groups):
        if i:
            par.add_run(',').font.superscript = True
        kind = KIND[key][1 if len(entries) > 1 else 0]
        lines = [f'{label}. {text}' if label else text for label, text in entries]
        lines[0] = f'{kind}. {lines[0]}'
        token = f'[[EN{len(ENDNOTES) + 1:03d}]]'
        ENDNOTES.append((token, '\r'.join(lines)))
        par.add_run(token)


def figure(doc, png, caption, width=None, height=None):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.left_indent = BODY_INDENT
    par.paragraph_format.keep_with_next = True
    par.add_run().add_picture(str(png), width=width, height=height)
    cap = doc.add_paragraph(style='Caption')
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.left_indent = BODY_INDENT
    cap.add_run(caption)


def landscape(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin = sec.right_margin = Inches(0.75)
    sec.top_margin = sec.bottom_margin = Inches(0.6)
    return sec


def portrait(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
    return sec


def md_runs(text):
    """**bold** markup in rubric-core.md to runs."""
    runs = []
    for i, part in enumerate(re.split(r'\*\*', text)):
        if part:
            runs.append((part, i % 2 == 1))
    return runs


def overview_blocks():
    """The 'Accessible Procurement Process Overview' section of rubric-core.md, as blocks.
    Its ### headings become subsections of the overview."""
    md = (HERE / 'rubric-core.md').read_text(encoding='utf-8')
    start = md.index('## Accessible Procurement Process Overview')
    end = md.index('## 1. Inputs and symbols')
    body = md[start:end].split('\n', 1)[1]
    blocks = []
    for chunk in re.split(r'\n\s*\n', body.strip()):
        prose, bullet = [], None
        for line in chunk.strip().split('\n'):
            if line.startswith('### '):
                blocks.append(('h3', [(line[4:].strip(), False)]))
                continue
            if line.startswith('- '):
                if prose:
                    blocks.append(('p', md_runs(' '.join(prose))))
                    prose = []
                bullet = [line[2:].strip()]
                blocks.append(('li', bullet))
            elif bullet is not None and line.startswith('  '):
                bullet.append(line.strip())      # a wrapped bullet continues on an indented line
            else:
                bullet = None
                prose.append(line.strip())
        if prose:
            blocks.append(('p', md_runs(' '.join(prose))))
    return [(k, md_runs(' '.join(v)) if k == 'li' else v) for k, v in blocks]


def plain(runs):
    """The runs of a paragraph with bold removed (the document rule). A short bold label that
    opens the paragraph and ends in a full stop, such as "Burden.", keeps its bold, since it is
    a labelled entry and not a key concept inside the text."""
    if runs and runs[0][1] and runs[0][0].rstrip().endswith('.') and len(runs[0][0]) <= 40:
        return [runs[0]] + [(text, False) for text, _ in runs[1:]]
    return [(text, False) for text, _ in runs]


def render(doc, blocks, heading=None, level=1, cite=True):
    """Renders the blocks. Headings take `level`; paragraphs, lists and tables sit in the body
    column, indented to where the heading text starts. The first paragraph after a heading
    carries that heading's endnote markers, so each note sits on the definition it grounds.
    With `cite` false, only the `heading` passed in is cited: the overview's subsections share
    names with definitions and must not repeat their notes."""
    pending = heading
    for kind, payload in blocks:
        if kind == 'h3':
            text = ''.join(t for t, _ in payload)
            doc.add_heading(text, level=level)
            pending = text if cite else None
        elif kind == 'p':
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = BODY_INDENT
            add_runs(par, plain(payload))
            if pending:
                mark(par, pending)
                pending = None
        elif kind == 'li':
            par = doc.add_paragraph(style='List Bullet')
            par.paragraph_format.left_indent = BODY_INDENT + Inches(0.25)
            par.paragraph_format.first_line_indent = Inches(-0.25)
            add_runs(par, payload)
        elif kind == 'table':
            add_table(doc, payload)


def body_paragraph(doc, text=''):
    """A paragraph in the body column."""
    par = doc.add_paragraph(text)
    par.paragraph_format.left_indent = BODY_INDENT
    return par


def unnumbered(par):
    """Drops the heading number from one heading paragraph (a numId of 0 means no list)."""
    num = par._p.get_or_add_pPr().get_or_add_numPr()
    ilvl = OxmlElement('w:ilvl')
    ilvl.set(qn('w:val'), '0')
    nid = OxmlElement('w:numId')
    nid.set(qn('w:val'), '0')
    num.append(ilvl)
    num.append(nid)
    return par


HEADER_FILL = 'E8EEF0'
ACCENT = RGBColor(0x0A, 0x56, 0x60)
GREY = RGBColor(0x44, 0x44, 0x44)
TEXT_WIDTH = Inches(6.5)          # letter page at 1 inch margins
BODY_INDENT = Inches(0.5)         # body text starts where the heading text starts
HEADING_NUM_ID = 90               # the numbering definition linked to the heading styles


def style(doc):
    st = doc.styles
    normal = st['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, before in (('Heading 1', 16, 18), ('Heading 2', 13, 12), ('Heading 3', 11, 8)):
        h = st[name]
        h.font.name = 'Calibri'
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.italic = False
        h.font.color.rgb = ACCENT
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        # Word reads the heading font from the theme unless the East Asian slot is set too.
        rpr = h.element.get_or_add_rPr()
        fonts = rpr.find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts'); rpr.append(fonts)
        for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
            fonts.set(qn(attr), 'Calibri')
        for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if fonts.get(qn(attr)) is not None:
                del fonts.attrib[qn(attr)]
    t = st['Title']
    t.font.name = 'Calibri'; t.font.size = Pt(26); t.font.bold = True; t.font.color.rgb = ACCENT
    lb = st['List Bullet']
    lb.paragraph_format.space_after = Pt(3)
    cap = st['Caption']
    cap.font.size = Pt(10); cap.font.italic = True; cap.font.bold = False
    cap.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    cap.paragraph_format.space_after = Pt(12)
    sec = doc.sections[0]
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, side, Inches(1))
    heading_numbering(doc)


def heading_numbering(doc):
    """Word's multilevel list linked to the heading styles: 1, 1.1, 1.1.1, with the number
    hanging in the margin so the heading text starts at the body column."""
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), str(HEADING_NUM_ID))
    mlt = OxmlElement('w:multiLevelType')
    mlt.set(qn('w:val'), 'multilevel')
    abstract.append(mlt)
    for ilvl, (style_id, text) in enumerate((('Heading1', '%1'), ('Heading2', '%1.%2'), ('Heading3', '%1.%2.%3'))):
        lvl = OxmlElement('w:lvl')
        lvl.set(qn('w:ilvl'), str(ilvl))
        for tag, val in (('w:start', '1'), ('w:numFmt', 'decimal'), ('w:pStyle', style_id), ('w:lvlText', text), ('w:lvlJc', 'left')):
            el = OxmlElement(tag)
            el.set(qn('w:val'), val)
            lvl.append(el)
        ppr = OxmlElement('w:pPr')
        ind = OxmlElement('w:ind')
        ind.set(qn('w:left'), str(int(BODY_INDENT.inches * 1440)))
        ind.set(qn('w:hanging'), str(int(BODY_INDENT.inches * 1440)))
        ppr.append(ind)
        lvl.append(ppr)
        abstract.append(lvl)
    nums = numbering.findall(qn('w:num'))
    if nums:
        nums[0].addprevious(abstract)
    else:
        numbering.append(abstract)
    num = OxmlElement('w:num')
    num.set(qn('w:numId'), str(HEADING_NUM_ID))
    ref = OxmlElement('w:abstractNumId')
    ref.set(qn('w:val'), str(HEADING_NUM_ID))
    num.append(ref)
    numbering.append(num)
    for ilvl, name in enumerate(('Heading 1', 'Heading 2', 'Heading 3')):
        num_pr = doc.styles[name].element.get_or_add_pPr().get_or_add_numPr()
        lvl = OxmlElement('w:ilvl')
        lvl.set(qn('w:val'), str(ilvl))
        nid = OxmlElement('w:numId')
        nid.set(qn('w:val'), str(HEADING_NUM_ID))
        num_pr.append(lvl)
        num_pr.append(nid)


def rule(par, side):
    """A thin rule under (side 'bottom') or over (side 'top') the paragraph."""
    ppr = par._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    line = OxmlElement(f'w:{side}')
    line.set(qn('w:val'), 'single')
    line.set(qn('w:sz'), '6')
    line.set(qn('w:space'), '4')
    line.set(qn('w:color'), 'BDBAB0')
    bdr.append(line)
    ppr.insert_element_before(bdr, 'w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap', 'w:overflowPunct',
                              'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd', 'w:snapToGrid',
                              'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
                              'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle',
                              'w:rPr', 'w:sectPr', 'w:pPrChange')


def small(run):
    run.font.size = Pt(9)
    run.font.color.rgb = GREY
    return run


def right_tab_only(par, width):
    """One right tab at the text edge. The Header and Footer styles carry a centre tab at 3.25
    inches and a right tab at 6.5, which only a clearing tab stop at the same position removes."""
    tabs = par.paragraph_format.tab_stops
    for pos in (Inches(3.25), Inches(6.5)):
        if pos != width:
            tabs.add_tab_stop(pos, WD_TAB_ALIGNMENT.CLEAR)
    tabs.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)


def furnish(doc, version):
    """The running header and footer of every section, sized to that section's text width.
    The cover page, the first page of the first section, carries neither."""
    for i, sec in enumerate(doc.sections):
        width = sec.page_width - sec.left_margin - sec.right_margin
        sec.different_first_page_header_footer = i == 0
        if i:
            sec.header.is_linked_to_previous = False
            sec.footer.is_linked_to_previous = False
        head = sec.header.paragraphs[0]
        head.text = ''
        right_tab_only(head, width)
        small(head.add_run('SFBRN Accessible Procurement Rubric'))
        small(head.add_run(f'\tMatrix version {version}, {date.today().isoformat()}'))
        rule(head, 'bottom')
        foot = sec.footer.paragraphs[0]
        foot.text = ''
        right_tab_only(foot, width)
        small(foot.add_run('San Francisco Bay Region Network'))
        small(foot.add_run('\tPage '))
        field(foot, 'PAGE', '1')
        small(foot.add_run(' of '))
        field(foot, 'NUMPAGES', '1')
        for r in foot.runs:
            small(r)
        rule(foot, 'top')


def update_fields_with_word(path):
    """Opens the saved document in Word, turns every endnote marker into a Word endnote, inserts
    the table of contents at its marker, updates the page fields, and saves. The contents page then carries the real page
    numbers and a link to each heading, and the notes sit at the end under their own heading."""
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx('Word.Application')
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(str(path), ReadOnly=False, AddToRecentFiles=False)
        try:
            if d.ReadOnly:
                raise RuntimeError('Word opened the file read-only, which means it repaired malformed XML. Check element order.')
            d.ActiveWindow.View.Type = 3          # wdPrintView; Read Mode refuses the endnote settings
            d.Endnotes.NumberStyle = 0            # wdNoteNumberStyleArabic
            d.Endnotes.Location = 1               # wdEndOfDocument
            d.Styles('Endnote Text').ParagraphFormat.SpaceAfter = 4   # air between the entries of one note
            for token, text in ENDNOTES:
                rng = d.Content
                found = rng.Find.Execute(FindText=token, MatchCase=True, MatchWildcards=False)
                if not found:
                    raise RuntimeError(f'endnote marker not found: {token}')
                rng.Text = ''
                d.Endnotes.Add(rng).Range.Text = text   # the Text argument of Add is ignored
            d.Repaginate()
            heading = d.Content
            if heading.Find.Execute(FindText='Contents', MatchCase=True, MatchWholeWord=True):
                heading.Paragraphs(1).Style = 'TOC Heading'
                heading.Paragraphs(1).Range.ListFormat.RemoveNumbers()
            rng = d.Content
            if not rng.Find.Execute(FindText='[[TOC]]', MatchCase=True, MatchWildcards=False):
                raise RuntimeError('contents marker not found')
            rng.Text = ''
            # Positional, since late-bound COM drops named arguments: Range, UseHeadingStyles,
            # UpperHeadingLevel, LowerHeadingLevel, UseFields, TableID, RightAlignPageNumbers,
            # IncludePageNumbers, AddedStyles, UseHyperlinks, HidePageNumbersInWeb, UseOutlineLevels.
            d.TablesOfContents.Add(rng, True, 1, 2, False, '', True, True, '', True, True, False)
            for toc in d.TablesOfContents:
                toc.Update()
            d.Fields.Update()
            d.Save()
        finally:
            d.Close(SaveChanges=0)
    finally:
        word.Quit()


def main():
    tmp = tempfile.mkdtemp()
    subprocess.run(['node', str(HERE / 'dump-definitions.js'), str(PLAYWRIGHT), tmp], check=True)
    data = json.loads((Path(tmp) / 'defs.json').read_text(encoding='utf-8'))
    parser = Defs()
    parser.feed(data['html'])

    doc = Document()
    style(doc)

    # ---- Title page and contents.
    for _ in range(6):
        doc.add_paragraph()
    doc.add_paragraph('SFBRN Accessible Procurement Rubric', style='Title')
    sub = doc.add_paragraph()
    sub.add_run('The accessibility review matrix of the San Francisco Bay Region Network.').font.size = Pt(14)
    meta = doc.add_paragraph()
    meta.add_run(f'Review matrix version {data["version"]}. Built {date.today().isoformat()} from review-matrix.js, '
                 'the model the SFBRN campuses run.').italic = True
    doc.add_paragraph()
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    doc.add_paragraph(
        'ADA Title II and Section 508, which Government Code 7405 and 11135 apply to the CSU, require the web pages, '
        'apps, software and hardware a campus provides to meet WCAG 2.1 AA or the Section 508 Standards. CSU '
        'policy has the campus Accessible Technology Initiative (ATI) reviewer check every technology purchase '
        'against that standard before the campus buys the product. The Chancellor\'s Office procurement process '
        'names the steps of that check and leaves the judgment inside each step to the campus. This document is '
        'the rubric by which the three San Francisco Bay Region Network (SFBRN) campuses, CSU East Bay, '
        'San Francisco State and Sonoma State, make that judgment.')
    doc.add_page_break()
    # Word builds the table of contents itself in the Word pass, at the [[TOC]] token.
    doc.add_paragraph('Contents')
    doc.add_paragraph('[[TOC]]')

    sections = ['Accessible Procurement Process Overview'] + [''.join(t for t, _ in p) for k, p in parser.blocks if k == 'h3']
    missing = [h for h in GROUNDING if h not in sections]
    if missing:
        sys.exit(f'grounding has no section: {missing}')

    # ---- Overview and the process figures.
    doc.add_page_break()
    doc.add_heading('Accessible Procurement Process Overview', level=1)
    render(doc, overview_blocks(), 'Accessible Procurement Process Overview', level=2, cite=False)
    figure(doc, Path(tmp) / 'review-process-checks.png',
           'Figure 1. The shared checks, from requisition intake to the risk grade. An accommodation purchase is '
           'approved at intake. Holds 1 to 3 pause the review, '
           'and hold 3 resumes at the Severity step. With no evidence at all, the review falls back once the ATI '
           'review team records that none can be obtained.', width=Inches(6.0))
    figure(doc, Path(tmp) / 'review-process-remedy.png',
           'Figure 2. The remedy step, from an Alternate access required risk grade to Deny or Approve with TAAP. '
           'Holds 4 and 5 pause the review. Under the fallback the minimum is provisional, and the first TAAP '
           'review re-rates Severity on evidence at the Veracity floor.', width=Inches(5.8))

    # ---- The rubric table, on a landscape page.
    landscape(doc)
    doc.add_heading('The rubric on one page', level=1)
    figure(doc, Path(tmp) / 'rubric-core.png',
           'Figure 3. The rubric core, generated from the same model as this document. It collects the values '
           'the sections that follow define, the risk grade for each pairing of Exposure and Severity, the '
           'Veracity floor, the TAAP standard by pressure and the warning level.',
           width=Inches(9.0))
    portrait(doc)

    # ---- The definitions, one numbered subsection each.
    doc.add_heading('Definitions', level=1)
    current = None
    section = []
    for kind, payload in parser.blocks:
        if kind == 'h3':
            if current:
                render(doc, section, level=2)
            current = ''.join(t for t, _ in payload)
            section = [(kind, payload)]
        else:
            section.append((kind, payload))
    if current:
        render(doc, section, level=2)

    # ---- Notes. Word places the endnotes after this heading.
    doc.add_page_break()
    unnumbered(doc.add_heading('Notes', level=1))
    body_paragraph(doc,
        'Each note names a rule that binds the CSU, a principle the Accessible Technology Initiative '
        'applies in every review, or a calibration choice that no rule sets and the pilot tests. '
        'Guidance, settlements and cases from other institutions are not cited, because they do not '
        'bind the CSU.')

    shutil.rmtree(tmp, ignore_errors=True)
    lock = OUT.with_name('~$' + OUT.name[2:])
    if lock.exists():
        sys.exit(f'{OUT.name} is open in Word. Close it and run again.')
    furnish(doc, data['version'])
    doc.save(OUT)
    update_fields_with_word(OUT)
    print('wrote', OUT)


if __name__ == '__main__':
    main()
