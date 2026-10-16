"""The annual report: payload reduction, assembly, and a rendered document. No database."""
from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from app.database.tools.report_export.annual_report import model as M
from app.database.tools.report_export.annual_report.docx_renderer import render

pytestmark = pytest.mark.unit


def _payload(key="1.1-ins", status="Defined", previous="Initiated", **over):
    base = {
        "indicator": {
            "composite_key": key, "success_indicator": "Campus has formally documented a process.",
            "number": key.rsplit("-", 1)[0], "removed": False, "goal_number": 1,
            "goal_name": "Establish accountability", "working_group": "Instructional Materials",
            "examples_of_evidence": "A policy.", "established_example": "A published procedure.",
        },
        "evidence_coverage": {"summary": {"total": 2, "satisfied": 1}, "requirements": [
            {"level": "Established", "element": "Procedures", "requirement": "A written procedure",
             "rubric_dimension": "procedures", "satisfied": True, "satisfied_by": [{"title": "Policy X"}]},
            {"level": "Established", "element": "Position", "requirement": "A named owner",
             "rubric_dimension": "resources", "satisfied": False, "satisfied_by": [],
             "implementation_evidenced": False},
        ]},
        "year": "2025-2026",
        "campus": {"abbreviation": "sfsu", "name": "San Francisco State University"},
        "status": {"status_level": status, "status_value": "2", "previous_status_level": previous},
        "yse": {"year_identifier": f"2025-2026-{key}-sfsu", "unique_id": "u1",
                "administrative_review_complete": True, "administrative_review_completed_date": "2026-06-01",
                "admin_review_description": "The campus has a policy and is rebuilding the practice.",
                "priority_level": "High", "will_work_on_next_year": True},
        "people": {"implementers": [{"name": "A. Person", "title": "Director", "roles": [
            {"name": "ATI Coordinator", "in_position_description": True}]}],
            "admin_review_completed_by": {"name": "R. Viewer"}},
        "recommendations": [{"recommendation": "Name the owner.", "detail": "In Academic Affairs.",
                             "status": "open", "date_created": "2026-05-01", "created_by": {"name": "D. F."}}],
        "concerns": [{"concern": "Process lapsed.", "status": "converted", "date_raised": "2026-02-01",
                      "raised_by": {"name": "D. F."}, "became": {"kind": "plan", "text": "Rebuild it."}}],
        "admin_review_notes": [{"content": "Hold at Defined.", "date_created": "2026-06-01",
                                "created_by": {"name": "R. Viewer"}}],
        "community_stakeholders": [{"name": "Bookstore Advisory Committee", "note": "owns the memo"}],
        "implementations": [
            {"type": "InternalPolicy", "title": "Policy X", "description": "The policy.",
             "strength": 3, "control": "internal", "rationale": "It is the policy.", "retired": False,
             "owner": {"name": "S. Owner", "title": "Dean"}, "accountable_working_group": {"name": "Instructional Materials"},
             "documents": [{"name": "Policy.pdf", "uri_path": "https://x/policy.pdf"}],
             "webpages": [{"name": "Policy page", "url": "https://x/policy"}],
             "participants": [{"person": {"name": "P. One"}, "role": "author"}]},
            {"type": "Process", "title": "Old process", "description": "", "retired": True,
             "retired_date": "2024-01-01", "retired_note": "Replaced.", "owner": None,
             "documents": [], "webpages": []},
        ],
        "taaps": [{"title": "TAAP 1"}],
        "assets": [{"title": "Canvas", "scope": "campus"}],
        "interfaces": [{"title": "Course shell", "function": "delivery"}],
        "tools": [{"title": "UDOIT"}],
        "vendors": [{"name": "Instructure"}],
        "plans": [{"name": "Rebuild it.", "description": "Memo cascade.", "plan_status": "In Progress",
                   "is_key_plan": True}],
        "accomplishments": [{"name": "Committee reformed", "description": "Met in December."}],
        "notes": [
            {"name": "Current state", "content": "Lapsed since 2019.", "dateCreated": "2024-12-11"},
            {"name": "Subcommittee", "content": "Draft going to the Provost.", "dateCreated": "2026-04-22"},
            {"name": "Old habit", "content": "Something undated.", "dateCreated": None},
        ],
        "messages": [{"name": "Meeting summary", "content": "Draft to Provost.", "date_sent": "2026-05-05",
                      "type": "meeting minutes"}],
        "metrics": [{"name": "Adoption rate", "metric_type": "percent", "single_value": "62"}],
    }
    base.update(over)
    return base


def test_reduce_indicator_flattens_the_payload():
    ind = M.reduce_indicator(_payload())
    assert ind["key"] == "1.1-ins" and ind["number"] == "1.1"
    assert ind["working_group"] == "Instructional Materials" and ind["working_group_abbrev"] == "ins"
    assert ind["status"] == {"current": "Defined", "current_value": "2", "previous": "Initiated"}
    assert ind["review"]["complete"] and ind["review"]["completed_by"] == "R. Viewer"
    assert ind["review"]["summary"].startswith("The campus has a policy")
    assert ind["implementers"][0]["roles"][0]["in_position_description"] is True
    assert [im["retired"] for im in ind["implementations"]] == [False, True]
    assert ind["implementations"][0]["documents"][0]["url"] == "https://x/policy.pdf"
    assert ind["implementations"][0]["participants"] == [{"name": "P. One", "role": "author"}]
    assert ind["plans"][0]["status"] == "In Progress" and ind["plans"][0]["is_key_plan"]
    assert ind["recommendations"][0]["author"] == "D. F."
    assert ind["concerns"][0]["became"] == "plan: Rebuild it."
    assert ind["messages"][0]["date"] == "2026-05-05"
    # Notes read the serializer's camelCase date and come newest first, undated last.
    assert [n["date"] for n in ind["notes"]] == ["2026-04-22", "2024-12-11", ""]
    assert ind["coverage"]["total"] == 2 and ind["coverage"]["satisfied"] == 1
    assert ind["coverage"]["requirements"][1]["implementation_evidenced"] is False
    assert ind["tools"] == ["UDOIT"] and ind["vendors"] == ["Instructure"]


def test_reduce_indicator_renders_list_properties_as_lines():
    p = _payload()
    p["indicator"]["examples_of_evidence"] = []
    p["indicator"]["established_example"] = ["A policy.", "", "A named owner."]
    ind = M.reduce_indicator(p)
    assert ind["guide"]["examples_of_evidence"] == ""
    assert ind["guide"]["established_example"] == "A policy.\nA named owner."


def test_reduce_indicator_tolerates_empty_payload():
    ind = M.reduce_indicator({"indicator": {"composite_key": "9.9-web"}})
    assert ind["key"] == "9.9-web" and ind["status"]["current"] == ""
    assert ind["implementations"] == [] and ind["coverage"]["total"] == 0


def test_indicator_sort_is_numeric():
    keys = ["1.11-web", "1.2-web", "1.1-web", "10.1-web", "2.3-web"]
    assert sorted(keys, key=M.indicator_sort_key) == ["1.1-web", "1.2-web", "1.11-web", "2.3-web", "10.1-web"]


def _model():
    inds = [
        M.reduce_indicator(_payload("1.1-ins")),
        M.reduce_indicator(_payload("1.11-ins", status="Established", previous="")),
        M.reduce_indicator(_payload("1.2-ins", status="", previous="")),
        M.reduce_indicator(_payload("3.1-web", status="Managed", implementations=[], plans=[],
                                    recommendations=[], concerns=[], notes=[], messages=[], metrics=[],
                                    assets=[], interfaces=[], tools=[], vendors=[], taaps=[],
                                    community_stakeholders=[], admin_review_notes=[],
                                    evidence_coverage={"summary": {"total": 0, "satisfied": 0}, "requirements": []})),
    ]
    inds[3]["working_group"] = "Web"
    inds[3]["working_group_abbrev"] = "web"
    inds[3]["goal_number"] = "3"
    inds[3]["goal_name"] = "Web goal"
    rubric = [{"value": str(i), "level": lvl, "procedures": "p", "resources": "r", "documentation": "d"}
              for i, lvl in enumerate(M.STATUS_ORDER)]
    return M.assemble("2025-2026", {"abbreviation": "sfsu", "name": "San Francisco State University"}, inds, rubric)


def test_assemble_groups_in_registry_order_and_counts_agree():
    model = _model()
    assert [wg["name"] for wg in model["working_groups"]] == ["Web", "Instructional Materials"]
    ins = model["working_groups"][1]
    assert [i["key"] for i in ins["goals"][0]["indicators"]] == ["1.1-ins", "1.2-ins", "1.11-ins"]
    s = model["summary"]
    assert s["indicators"] == 4 and s["working_groups"] == 2 and s["goals"] == 2
    assert s["established_or_above"] == 2
    assert s["status_distribution"]["Instructional Materials"]["Unset"] == 1
    assert s["evidence_items"] == 6 and s["active_evidence_items"] == 3
    assert s["open_recommendations"] == 3 and s["open_concerns"] == 0


def test_render_writes_a_valid_docx(tmp_path):
    model = _model()
    path = render(model, str(tmp_path / "report.docx"))
    assert Path(path).stat().st_size > 10_000
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
        assert "Annual Status Report" in xml
        assert "San Francisco State University" in xml
        assert "Executive summary" in xml
        assert "Instructional Materials" in xml and "Web" in xml
        assert "Campus has formally documented a process." in xml
        assert "The campus has a policy and is rebuilding the practice." in xml
        assert "Policy X" in xml and "Retired 2024-01-01" in xml
        assert "Name the owner." in xml
        assert "Status level definitions" in xml
        assert "TOC" in xml
        assert "No evidence is listed" in xml   # the empty Web indicator
        # Restructured per the 2026-09-29 review of the first draft.
        assert "Draft to Provost" not in xml            # correspondence is not reproduced
        assert "What the indicator asks for" not in xml  # the guide block is gone
        assert "Supplement: indicator bar coverage" in xml
        assert "Administrative review" in xml and "Hold at Defined." in xml
        assert "Bookstore Advisory Committee" in xml     # communities sit in the header
        assert 'w:name="cov_1_1_ins"' in xml and 'w:name="ind_1_1_ins"' in xml
        assert 'w:anchor="cov_1_1_ins"' in xml and 'w:anchor="ind_1_1_ins"' in xml
        assert "1 of 2 met" in xml and "No bar authored" in xml
        assert "Indicators without an authored bar" in xml
        # People come before the evidence table on the page, as bullets.
        assert xml.index("Responsible") < xml.index("Policy X")
        assert "Undated" in xml   # a note the record never dated says so
        assert xml.index("2026-04-22 · Subcommittee") < xml.index("2024-12-11 · Current state") < xml.index("Undated · Old habit")
        assert "word/media/image1.png" in z.namelist()   # the cover logo
        styles = z.read("word/styles.xml").decode("utf-8")
        assert "Report Kicker" in styles and "Georgia" in styles
    d = __import__("docx").Document(path)
    bullets = [p.text for p in d.paragraphs if p.style.name == "List Bullet"]
    assert any(b.startswith("A. Person, Director") for b in bullets)
    assert any(b.startswith("Bookstore Advisory Committee") for b in bullets)


def test_each_indicator_starts_on_its_own_page(tmp_path):
    import docx as _docx
    model = _model()
    d = _docx.Document(render(model, str(tmp_path / "pages.docx")))
    heads = [(p.text, bool(p.paragraph_format.page_break_before))
             for p in d.paragraphs if p.style.name in ("Heading 1", "Heading 2", "Heading 3")]

    def breaks(prefix):
        return next(b for text, b in heads if text.startswith(prefix))

    # Instructional Materials goal 1 holds 1.1, 1.2, 1.11 in that order: the goal heading
    # breaks, the first indicator rides with it, the rest break.
    assert breaks("Goal 1  Establish") is True
    assert breaks("1.1  ") is False
    assert breaks("1.2  ") is True
    assert breaks("1.11  ") is True
    assert breaks("2  Web") is True and breaks("3  Instructional") is True
    assert breaks("Appendix A") is True and breaks("Appendix B") is True


def test_render_is_deterministic_for_the_same_model(tmp_path):
    model = _model()
    a = render(model, str(tmp_path / "a.docx"))
    b = render(model, str(tmp_path / "b.docx"))
    with zipfile.ZipFile(a) as za, zipfile.ZipFile(b) as zb:
        assert za.read("word/document.xml") == zb.read("word/document.xml")
